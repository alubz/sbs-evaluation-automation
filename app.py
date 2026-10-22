import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from functools import lru_cache
from textwrap import wrap

import matplotlib
matplotlib.use("Agg")  # headless backend: no GUI, faster, thread/server safe
from matplotlib.figure import Figure
import numpy as np
import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Constants
CORE_COLUMNS = ['Program Name', 'Faculty Name', 'Session', 'Category', 'Comments']
TEXT_KEY_COLUMNS = ['Program Name', 'Faculty Name', 'Session']
EXCEL_SHEET = "Comments & Scores"
LOGO_CANDIDATES = ("logo.png", "input_file_0.png", "SBS_Logo.png")
SCORE_MAX = 5
EXCEL_BLUE = "#4472C4"
EXCEL_RED = "#C00000"  # Professional Excel dark red for charts
CHART_DPI = 150

def get_evaluation_params(df):
    """Dynamically detect evaluation parameter columns from the dataframe, excluding metadata and year columns."""
    return [
        col for col in df.columns 
        if col not in CORE_COLUMNS 
        and 'year' not in col.lower() 
        and pd.api.types.is_numeric_dtype(df[col])
    ]

def validate_dataframe(df):
    """Validate that required core columns exist in the dataframe."""
    missing_cols = [col for col in CORE_COLUMNS if col not in df.columns]
    if missing_cols:
        return False, f"Missing required columns: {', '.join(missing_cols)}"
    return True, None

def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Normalise headers and key text columns so groups don't split on stray spaces."""
    df = df.dropna(how="all").copy()
    df.columns = df.columns.astype(str).str.strip()
    for col in TEXT_KEY_COLUMNS:
        if col in df.columns:
            s = df[col]
            df[col] = s.where(s.isna(), s.astype(str).str.strip()).replace("", np.nan)
    return df

@st.cache_data(show_spinner="Reading file…")
def load_uploaded_file(file_bytes: bytes, name: str) -> pd.DataFrame:
    buf = io.BytesIO(file_bytes)
    if name.lower().endswith(".csv"):
        return clean_dataframe(pd.read_csv(buf))
    xls = pd.ExcelFile(buf)
    if EXCEL_SHEET not in xls.sheet_names:
        raise ValueError(f"Sheet '{EXCEL_SHEET}' not found. Sheets in file: {', '.join(xls.sheet_names)}")
    return clean_dataframe(xls.parse(EXCEL_SHEET, header=1))

@st.cache_data(show_spinner="Fetching Google Sheet…", ttl=300)
def load_google_sheet(sheet_url: str):
    match = re.search(r"/d/([a-zA-Z0-9-_]+)", sheet_url)
    if not match:
        return None, "Invalid Google Sheets URL format."
    gid_match = re.search(r"gid=([0-9]+)", sheet_url)
    gid = gid_match.group(1) if gid_match else "0"
    export_url = f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid}"
    try:
        return clean_dataframe(pd.read_csv(export_url)), None
    except Exception as e:
        return None, str(e)

# Dataclasses for structured data
@dataclass
class FacultyResult:
    session: str
    faculty: str
    respondents: int
    params: list
    scores: list
    prog_avgs: list
    avg_score: float
    avg_prog: float
    comments: list
    chart: bytes = b""

@dataclass
class DeptResult:
    session: str
    params: list
    scores: list
    avg_score: float
    comments: list
    chart: bytes = b""

@dataclass
class ProgrammeResult:
    faculty: list = field(default_factory=list)
    dept: list = field(default_factory=list)
    highest_session_respondents: int = 0

# Helper functions
def _to_pct(x):
    return (x / SCORE_MAX * 100).round(0)

# Analysis functions
def analyse_programme(prog_df: pd.DataFrame, param_cols: list) -> ProgrammeResult:
    result = ProgrammeResult()
    fac_rows = prog_df[prog_df["Faculty Name"].notna()]
    other_rows = prog_df[prog_df["Faculty Name"].isna()]

    if not fac_rows.empty:
        grouped = fac_rows.groupby(["Session", "Faculty Name"])
        scores = _to_pct(grouped[param_cols].mean())
        prog_avg = scores.mean().round(0)
        score_rows = scores.to_dict("index")
        result.highest_session_respondents = int(fac_rows.groupby("Session").size().max())

        for (session, faculty), group in grouped:
            row = score_rows[(session, faculty)]
            params = [p for p in param_cols if pd.notna(row[p])]
            ind = [row[p] for p in params]
            pav = [prog_avg[p] for p in params]
            result.faculty.append(FacultyResult(
                session=str(session), faculty=str(faculty), respondents=len(group),
                params=params, scores=ind, prog_avgs=pav,
                avg_score=round(sum(ind) / len(ind), 0) if ind else 0,
                avg_prog=round(sum(pav) / len(pav), 0) if pav else 0,
                comments=group["Comments"].dropna().tolist(),
            ))

    for session, group in other_rows.groupby("Session"):
        params = [p for p in param_cols if group[p].notna().any()]
        ind = _to_pct(group[params].mean()).tolist() if params else []
        result.dept.append(DeptResult(
            session=str(session), params=params, scores=ind,
            avg_score=round(sum(ind) / len(ind), 0) if ind else 0,
            comments=group["Comments"].dropna().tolist(),
        ))
    return result

# Chart rendering (memory-based)
def render_chart(params, series, title, series_labels, colors=None) -> bytes:
    """Excel-style clustered column chart. Returns PNG bytes."""
    colors = colors or [EXCEL_BLUE, EXCEL_RED]
    n = len(series)

    fig = Figure(figsize=(12, 6.5), dpi=CHART_DPI)
    fig.patch.set_facecolor("white")
    ax = fig.subplots()

    x = np.arange(len(params))
    bar_w = 0.5 / n
    
    valid_series = [np.asarray(s, dtype=float) for s in series if len(s) > 0]
    if valid_series:
        all_vals = np.concatenate(valid_series)
        min_val = int(np.nanmin(all_vals)) if not np.isnan(all_vals).all() else 0
        ymin = max(0, min_val - 4)
    else:
        ymin = 0

    for i, scores in enumerate(series):
        offset = (i - (n - 1) / 2) * bar_w
        bars = ax.bar(x + offset, scores, bar_w, color=colors[i % len(colors)], label=series_labels[i], edgecolor='white', linewidth=1)
        
        formatted_labels = [
            f"{int(round(v))}%" if (pd.notna(v) and not np.isnan(v)) else "0%" 
            for v in scores
        ]
        ax.bar_label(bars, labels=formatted_labels, padding=4, fontsize=9, fontweight='bold')

    ax.set_ylim(0, 135)
    ax.set_yticks(range(0, 140, 10))
    ax.tick_params(axis="y", labelsize=11)

    wrap_width = max(12, min(22, int(140 / max(len(params), 1))))
    wrapped_params = ["\n".join(wrap(str(p), wrap_width)) for p in params]

    ax.set_xticks(x)
    ax.set_xticklabels(wrapped_params, fontsize=10, rotation=45, ha='right', rotation_mode="anchor")
    ax.tick_params(axis="x", length=0)

    ax.set_ylabel("% Score", fontsize=12, fontweight="bold")
    ax.set_xlabel("Evaluation Parameters", fontsize=12, fontweight="bold", labelpad=10)
    ax.set_title(title, fontsize=13, fontweight="bold", pad=15)

    ax.grid(axis='y', alpha=0.3, linestyle='--')
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_linewidth(0.5)

    handles, labels = ax.get_legend_handles_labels()
    ax.legend(handles, labels, loc="upper right", fontsize=11, framealpha=0.9)
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=CHART_DPI, bbox_inches="tight", pad_inches=0.4)
    return buf.getvalue()

# Word document helpers
@lru_cache(maxsize=1)
def _logo_bytes():
    for fname in LOGO_CANDIDATES:
        if os.path.exists(fname):
            with open(fname, "rb") as f:
                return f.read()
    return None

def new_doc() -> Document:
    doc = Document()
    logo = _logo_bytes()
    if logo:
        hp = doc.sections[0].header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        try:
            hp.add_run().add_picture(io.BytesIO(logo), width=Inches(1.5))
        except Exception:
            pass
    return doc

def doc_to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def safe_name(text, max_len=40) -> str:
    return re.sub(r"[\\/:*?\"<>|'()\r\n]+", "_", str(text)).strip()[:max_len] or "unnamed"

def add_title(doc, text, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)

def _fill(cell, text, bold=False, size=9, center=False):
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Enhanced Professional Corporate Table Styling
def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, start=150, end=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', start), ('right', end)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_score_table(doc, params, scores, avg, prog_avgs=None, avg_prog=None, avg_label="Average scores"):
    compare = prog_avgs is not None
    headers = ["Evaluation Parameter", "(%) Individual Score" if compare else "(%) Score"]
    if compare:
        headers.append("(%) Programme Average")

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header Row Styling
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        set_cell_background(cell, "1F4E78")  # Dark Executive Blue
        set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(headers[i])
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows with Zebra Striping
    for idx, name in enumerate(params):
        row_cells = table.add_row().cells
        bg_color = "F9FAFC" if idx % 2 == 1 else "FFFFFF"
        
        for cell in row_cells:
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, start=180, end=180)

        _fill(row_cells[0], name, size=9, center=False)
        s_val = f"{int(round(scores[idx]))}%" if pd.notna(scores[idx]) else "N/A"
        _fill(row_cells[1], s_val, center=True, size=9)

        if compare:
            p_val = f"{int(round(prog_avgs[idx]))}%" if pd.notna(prog_avgs[idx]) else "N/A"
            _fill(row_cells[2], p_val, center=True, size=9)

    # Summary / Footer Row
    total_cells = table.add_row().cells
    for cell in total_cells:
        set_cell_background(cell, "EBF1F5")
        set_cell_margins(cell, top=120, bottom=120, start=180, end=180)

    _fill(total_cells[0], avg_label, bold=True, size=9.5)
    avg_val = f"{int(round(avg))}%" if pd.notna(avg) else "0%"
    _fill(total_cells[1], avg_val, bold=True, center=True, size=9.5)

    if compare:
        avg_prog_val = f"{int(round(avg_prog))}%" if pd.notna(avg_prog) else "0%"
        _fill(total_cells[2], avg_prog_val, bold=True, center=True, size=9.5)

def add_comments(doc, comments, level):
    doc.add_heading("Comments & Feedback", level=level)
    if comments:
        for c in comments:
            doc.add_paragraph(str(c), style="List Bullet")
    else:
        doc.add_paragraph("No comments recorded.")

def add_signature(doc, recipient_name=None, recipient_title=None):
    doc.add_paragraph("\nSUBMITTED THROUGH:")
    doc.add_paragraph("SBS Evaluation Automation System")
    doc.add_paragraph("\nName: __________________________")
    doc.add_paragraph("Title: __________________________")
    doc.add_paragraph("Sign: __________________________")
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")

def add_faculty_content(doc, fr: FacultyResult, pic_width, comment_level):
    if fr.params:
        add_score_table(doc, fr.params, fr.scores, fr.avg_score, fr.prog_avgs, fr.avg_prog)
        doc.add_paragraph()
        doc.add_picture(io.BytesIO(fr.chart), width=Inches(pic_width))
    elif not fr.comments:
        doc.add_paragraph("No evaluation data available for this session.")
    add_comments(doc, fr.comments, comment_level)

# Report builders
def build_faculty_report(programme, fr: FacultyResult, session_participants) -> bytes:
    doc = new_doc()
    add_title(doc, "FACULTY EVALUATION")
    doc.add_paragraph(f"Programme Name: {programme}")
    doc.add_paragraph(f"Session: {fr.session}")
    doc.add_paragraph(f"Faculty Name: {fr.faculty}")
    doc.add_paragraph(f"No. of Respondents: {fr.respondents}")
    doc.add_paragraph(f"No. of Participants: {session_participants}")
    add_faculty_content(doc, fr, 6.5, 2)
    add_signature(doc)
    return doc_to_bytes(doc)

def build_overall_report(programme, res: ProgrammeResult, total_participants) -> bytes:
    doc = new_doc()
    highest = res.highest_session_respondents
    pct_resp = round(highest / total_participants * 100) if total_participants > 0 else 0

    add_title(doc, f"OVERALL PROGRAMME EVALUATION REPORT\n{programme}")
    doc.add_paragraph(f"No. of Respondents: {highest}")
    doc.add_paragraph(f"No. of Participants: {total_participants}")
    doc.add_paragraph(f"Response Rate: {pct_resp}%")
    doc.add_paragraph("-" * 50)

    if res.faculty:
        doc.add_heading("Faculty Reports", level=1)
        for fr in res.faculty:
            doc.add_heading(f"{fr.faculty} - {fr.session}", level=2)
            add_faculty_content(doc, fr, 6.5, 3)
            doc.add_page_break()

    if res.dept:
        doc.add_heading("Department Reports", level=1)
        for d in res.dept:
            doc.add_heading(d.session, level=2)
            if d.params:
                add_score_table(doc, d.params, d.scores, d.avg_score, avg_label="Average score")
                doc.add_paragraph()
                doc.add_picture(io.BytesIO(d.chart), width=Inches(6.5))
            add_comments(doc, d.comments, 3)
            doc.add_page_break()

    add_signature(doc)
    return doc_to_bytes(doc)

def build_master_report(df: pd.DataFrame, param_cols: list) -> bytes:
    fac = df[df["Faculty Name"].notna()]
    doc = new_doc()
    add_title(doc, "STRATHMORE BUSINESS SCHOOL\nMASTER DEPARTMENTAL EVALUATION REPORT")
    doc.add_paragraph(f"Total Department Records Analyzed: {len(df)}")
    doc.add_paragraph(f"Total Active Programs: {df['Program Name'].nunique()}")
    doc.add_paragraph("-" * 50)

    doc.add_heading("1. Program-Level Benchmarks", level=1)
    if not fac.empty:
        prog_score = (fac.groupby("Program Name")[param_cols].mean().mean(axis=1) / SCORE_MAX * 100).round(1)
        valid_prog_score = prog_score.dropna()
        if not valid_prog_score.empty:
            n_sessions = fac.drop_duplicates(["Program Name", "Session", "Faculty Name"]).groupby("Program Name").size()
            n_resp = fac.groupby("Program Name").size()

            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, h in zip(table.rows[0].cells, ["Programme", "Overall Average (%)", "Faculty Sessions", "Responses"]):
                _fill(cell, h, bold=True, size=10, center=True)
            for prog, score in prog_score.items():
                if pd.notna(score):
                    cells = table.add_row().cells
                    _fill(cells[0], prog)
                    _fill(cells[1], f"{score}", center=True)
                    _fill(cells[2], int(n_sessions[prog]), center=True)
                    _fill(cells[3], int(n_resp[prog]), center=True)
        else:
            doc.add_paragraph("No numeric evaluation data available for program-level benchmarks.")

        doc.add_heading("2. Departmental Parameter Averages (%)", level=1)
        glob = _to_pct(fac[param_cols].mean()).dropna()
        if not glob.empty:
            add_score_table(doc, glob.index.tolist(), glob.tolist(), round(glob.mean(), 0), avg_label="Overall average")
            doc.add_paragraph()
            png = render_chart(glob.index.tolist(), [glob.tolist()], "All Programmes : Departmental Parameter Averages",
                               ["Department Average"], [EXCEL_BLUE])
            doc.add_picture(io.BytesIO(png), width=Inches(6.5))
        else:
            doc.add_paragraph("No numeric evaluation data available for departmental averages.")
    else:
        doc.add_paragraph("No faculty evaluation records found.")

    add_signature(doc)
    return doc_to_bytes(doc)

# Batch generation (in-memory, guaranteed valid ZIP structure)
def generate_programme_zip(df, programmes, participants, totals, gen_overall, gen_individual,
                           param_cols, on_progress=None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used_folders = set()
        for idx, prog in enumerate(programmes):
            prog_df = df[df["Program Name"] == prog]
            res = analyse_programme(prog_df, param_cols)
            folder = _unique(safe_name(prog, 80), used_folders)

            # Render charts once, share between reports
            for fr in res.faculty:
                if fr.params:
                    fr.chart = render_chart(fr.params, [fr.scores, fr.prog_avgs],
                                            f"{fr.session} : {fr.faculty}", ["Individual", "Programme Avg"])
            if gen_overall:
                for d in res.dept:
                    if d.params:
                        d.chart = render_chart(d.params, [d.scores], f"{d.session} : Department Evaluation",
                                               ["Department Score"], [EXCEL_BLUE])

            if gen_individual:
                used_files = set()
                for fr in res.faculty:
                    key = f"{prog} -- {fr.faculty} -- {fr.session}"
                    session_headcount = participants.get(key, fr.respondents)
                    data = build_faculty_report(prog, fr, session_headcount)
                    fname = _unique(f"{safe_name(fr.faculty)}_{safe_name(fr.session)}.docx", used_files)
                    zf.writestr(f"{folder}/Faculty_Reports/{fname}", data)

            if gen_overall:
                total_headcount = totals.get(prog, res.highest_session_respondents)
                data = build_overall_report(prog, res, total_headcount)
                zf.writestr(f"{folder}/{folder}_Overall_Programme_Report.docx", data)

            if on_progress:
                on_progress((idx + 1) / len(programmes), prog)
    
    # Ensure stream pointer is reset to start
    buf.seek(0)
    return buf.getvalue()

def _unique(name, used: set) -> str:
    base, ext = os.path.splitext(name)
    candidate, i = name, 2
    while candidate in used:
        candidate = f"{base}_{i}{ext}"
        i += 1
    used.add(candidate)
    return candidate

# UI
def show_preview(df):
    with st.expander("📋 Preview Data"):
        st.dataframe(df.head(10))
        st.caption(f"Showing first 10 of {len(df)} rows")

def ingest():
    source = st.radio("Select Data Source:", ["Upload Excel / CSV File", "Google Sheets Public Link"])
    df, ok_msg = None, ""

    if source == "Upload Excel / CSV File":
        up = st.file_uploader("Upload Evaluation Dataset", type=["xlsx", "csv"])
        if up is None:
            return None
        try:
            df = load_uploaded_file(up.getvalue(), up.name)
        except Exception as e:
            st.error(f"Error reading file: {e}")
            return None
        ok_msg = f"Successfully loaded file with {len(df)} total rows."
    else:
        url = st.text_input("Paste Public Google Sheets Link (Anyone with link can view):")
        if not url:
            return None
        df, err = load_google_sheet(url)
        if err:
            st.error(err)
            return None
        ok_msg = f"Successfully connected to Google Sheet ({len(df)} rows loaded)."

    valid, msg = validate_dataframe(df)
    if not valid:
        st.error(msg)
        st.warning("Available columns: " + ", ".join(df.columns.tolist()))
        return None
    st.success(ok_msg)
    show_preview(df)
    return df

def single_programme_mode(df, param_cols):
    programmes = sorted(df["Program Name"].dropna().unique().tolist())
    selected = st.multiselect("Select Programmes:", options=programmes, default=[])
    if not selected:
        st.warning("Please select at least one programme.")
        return

    st.markdown("### 📝 Enter Participant Headcount Setup")
    
    participants = {}
    totals = {}

    for prog in selected:
        prog_df = df[df["Program Name"] == prog]
        counts = prog_df[prog_df["Faculty Name"].notna()].groupby(["Session", "Faculty Name"]).size()
        st.markdown(f"--- \n#### {prog}")
        
        same_count = st.checkbox(
            f"Same total participant count for all sessions in **{prog}**", 
            value=True, 
            key=f"same_count_{prog}"
        )
        
        if same_count:
            single_val = st.number_input(
                f"Participant count for all sessions & overall total in **{prog}**", 
                value=max(len(prog_df), 25), 
                step=1, 
                key=f"single_val_{prog}"
            )
            totals[prog] = single_val
            for (session, faculty), _ in counts.items():
                participants[f"{prog} -- {faculty} -- {session}"] = single_val
        else:
            st.markdown("Enter individual class sizes per session:")
            session_vals = []
            for (session, faculty), n in counts.items():
                val = st.number_input(
                    f"Participants for **{session}** ({faculty}) — [{n} responses in file]", 
                    value=max(int(n), 10), 
                    step=1, 
                    key=f"{prog} -- {faculty} -- {session}"
                )
                participants[f"{prog} -- {faculty} -- {session}"] = val
                session_vals.append(val)
            
            tot_val = st.number_input(
                f"Total Overall Participants for **{prog}**", 
                value=max(sum(session_vals), len(prog_df)), 
                step=1, 
                key=f"total_{prog}"
            )
            totals[prog] = tot_val

    st.markdown("---")
    gen_overall = st.checkbox("Generate Overall Programme Report", value=True)
    gen_individual = st.checkbox("Generate Individual Faculty Reports (ZIP)", value=True)
    
    if st.button("🚀 Generate Programme Reports", type="primary"):
        if not (gen_overall or gen_individual):
            st.warning("Tick at least one report type.")
        else:
            bar = st.progress(0.0, text="Starting…")
            data = generate_programme_zip(
                df, selected, participants, totals, gen_overall, gen_individual, param_cols,
                on_progress=lambda frac, name: bar.progress(frac, text=f"Finished {name}"),
            )
            bar.empty()
            st.session_state["prog_zip"] = (data, len(selected))

    # Result survives the rerun triggered by clicking download
    if "prog_zip" in st.session_state:
        data, n = st.session_state["prog_zip"]
        st.success(f"Generated reports for {n} programme(s) successfully!")
        st.download_button("📥 Download All Reports ZIP", data, file_name="Multiple_Programmes_Reports.zip",
                           mime="application/zip")

def master_mode(df, param_cols):
    st.markdown("### 🏢 Master Departmental Report (All Programmes Roll-up)")
    st.markdown("This will aggregate all programmes in the dataset into a comprehensive periodic departmental review.")
    if st.button("🚀 Generate Master Departmental Report", type="primary"):
        with st.spinner("Compiling departmental report across all programmes..."):
            st.session_state["dept_docx"] = build_master_report(df, param_cols)
    if "dept_docx" in st.session_state:
        st.success("Successfully generated Master Departmental Report!")
        st.download_button(
            "📥 Download Master Departmental Report (.docx)", st.session_state["dept_docx"],
            file_name="Master_Departmental_Evaluation_Report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

def main():
    st.set_page_config(page_title="SBS Evaluation Automation", layout="centered", page_icon="📊")
    st.title("📊 Strathmore Business School - Evaluation Automation")
    st.markdown("Generate Individual Faculty Reports, Programme Reports, and Periodic Departmental Reports.")

    df = ingest()
    if df is None:
        return

    param_cols = get_evaluation_params(df)
    if not param_cols:
        st.error("No evaluation parameters found in the data. Please ensure there are numeric columns for scoring.")
        return

    st.markdown("---")
    scope = st.selectbox("Select Report Generation Mode:", [
        "Single Programme Mode (Individual Faculty & Programme Reports)",
        "Master Departmental Report Mode (Cross-Programme Roll-up)",
    ])
    
    if "Single Programme Mode" in scope:
        single_programme_mode(df, param_cols)
    else:
        master_mode(df, param_cols)

if __name__ == "__main__":
    main()